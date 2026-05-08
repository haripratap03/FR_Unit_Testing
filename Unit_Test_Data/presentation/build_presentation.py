"""Generate the FR_Thor unit-testing PowerPoint deck.

Onboarding-friendly: every one of the 130 test cases is shown with a plain-
English description and a "why it matters" line so a new joinee can read the
deck end-to-end and understand both *what* the suite tests and *why*.

Run:
    python Unit_Test_Data/presentation/build_presentation.py
"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------------------
# Paths & data
# ----------------------------------------------------------------------------
HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent  # FR_Thor/
STATE_JSON = ROOT / "Unit_Test_Data" / "reports" / "test_state_20260505_120832.json"
OUT_PATH = HERE / "FR_Thor_Unit_Testing.pptx"

with STATE_JSON.open(encoding="utf-8") as fh:
    STATE = json.load(fh)


def find_test(tid: str) -> dict:
    """Look up a test entry by ID across all features."""
    if tid.startswith("FD"):
        return STATE["features"]["face_detection"]["tests"][tid]
    if tid.startswith("FT"):
        return STATE["features"]["face_tracking"]["tests"][tid]
    if tid.startswith("FR"):
        return STATE["features"]["face_recognition"]["tests"][tid]
    raise KeyError(tid)


# ----------------------------------------------------------------------------
# Per-test rationale — "why this test exists"
# Each line is plain English and meant to make sense to someone reading the
# deck without ever opening the source code.
# ----------------------------------------------------------------------------
RATIONALE: dict[str, str] = {
    # ---------- Face Detection — score_face ----------
    "FD-01": "Without this gate, every wall, shadow or poster YOLO mis-flags would be saved as a face and pollute downstream FR.",
    "FD-02": "Confirms the threshold is strict (<), not inverted; protects against off-by-one mistakes when tuning the conf knob.",
    "FD-03": "Sub-32 px crops lack the detail OSNet/InsightFace need; saving them creates noise embeddings and false matches.",
    "FD-04": "Boundary check that valid small faces (distant customers) still pass; one pixel tighter would lose them.",
    "FD-05": "Anchors the upper end of the score range. If the cleanest possible face only scores 30, the whole grading scheme is broken.",
    "FD-06": "Blur is the #1 reason a crop is unusable for FR; the blur sub-score must visibly drop on blurry inputs.",
    "FD-07": "Side-profile crops produce poor embeddings; the pose sub-score must downgrade them so frontal beats profile.",
    "FD-08": "Both ears visible is a strong frontal-pose signal; without it, a perfectly straight-ahead face could be wrongly tagged as turned.",
    "FD-09": "If YOLO's pose head misfires, score_face must still produce a usable score using a neutral fallback instead of crashing.",
    "FD-10": "Some YOLO weights ship without a pose head at all; the function must handle 'no keypoints' silently or every clip dies.",
    "FD-11": "Sun-glare frames are common at front-of-store cameras; overexposed shots must not score high just because they are sharp.",
    # ---------- Face Detection — check_quality ----------
    "FD-12": "Confirms the PASS path of the q-tag gate; if this breaks, every face is rejected and FR has nothing to work with.",
    "FD-13": "Confirms the FAIL path; without it, low-quality faces flood the FR stage and pollute customer clusters.",
    "FD-14": "Legacy files lack a q-tag; rejecting them silently would lose data, so the gate must default to allow.",
    "FD-15": "Boundary check that the floor is strict-less-than; if it slipped, q=25 faces would be discarded for no reason.",
    # ---------- Face Detection — skin-tone (public + private) ----------
    "FD-16": "Establishes the green light: if this fails, every face crop fails the skin gate and FR collapses to zero matches.",
    "FD-17": "Most YOLO false positives are walls, signs, jackets — non-skin colour is the cheapest way to drop them early.",
    "FD-18": "Confirms the centre-60% rule. A real face has skin in the middle of the crop, not just at the edges.",
    "FD-22": "Twin of FD-16 but on the actual pipeline.py helper; protects against the public/private wrappers diverging.",
    "FD-23": "Twin of FD-17 on the FAIL path; ensures the internal helper agrees with the test-facing wrapper.",
    # ---------- Face Detection — keypoints ----------
    "FD-19": "If InsightFace landmarks fall inside the YOLO box, the crop is genuinely face-shaped; this is the 3rd filter that drops garbage.",
    "FD-20": "Boundary on the min_keypoints=3 rule; if it slipped to 2, scarves, hats and half-faces would pass.",
    "FD-21": "If InsightFace cannot find a face inside what YOLO calls a face, the YOLO box is wrong — must fail closed.",
    "FD-24": "Sanity baseline for the bbox-containment counter; without this, every other counter test is meaningless.",
    "FD-25": "Negative case: counter must return 0, not silently wrap or treat 'outside' as 'unknown'.",
    "FD-26": "Boundary inclusivity — if the margin used < instead of <=, valid frontier faces would be dropped.",
    "FD-27": "Mixed case proves the counter is per-keypoint, not all-or-nothing.",
    # ---------- Face Detection — _enhance_contrast ----------
    "FD-28": "CLAHE is one of the four FR rescue strategies; if it doesn't actually brighten dark crops, the FR retry loop has no value.",
    "FD-29": "FR retries the same crop with multiple strategies; if one stage mutates the input, later stages get a tainted array.",
    # ---------- Face Detection — _detect_face ----------
    "FD-30": "When YOLO returns two faces in one crop (e.g. reflection), we must pick the dominant subject, not whichever is first.",
    "FD-31": "FAISS treats inner-product as cosine only on unit vectors; a non-normalised embedding silently breaks staff matching.",
    "FD-32": "A zero vector would 'match' nothing, but in cosine space it can create ghost matches; returning None is safer.",
    # ---------- Face Detection — _extract_with_strategies ----------
    "FD-33": "Validates the happy path of the 4-strategy ladder; if 'direct' never works, the whole ladder is wasted compute.",
    "FD-34": "Distant customers produce small crops; the upscale rung must trigger automatically for crops below the size floor.",
    "FD-35": "When 'direct' fails on a normal-size crop, the cause is usually contrast — second rung must be CLAHE, not upscale.",
    "FD-36": "Some borderline crops only embed once surrounded by black padding; this is the last hope before giving up.",
    "FD-37": "Returning (None, None) lets FR mark the person no_face cleanly; raising would crash the entire run.",
    # ---------- Face Detection — _batch_capture_faces ----------
    "FD-38": "Baseline: capture must actually save when nothing is wrong, otherwise the pipeline produces zero face files.",
    "FD-39": "Without the cooldown, every frame of a stationary customer would be saved as a separate file (tens of thousands).",
    "FD-40": "Defence in depth: even if score_face leaks a bad crop, capture re-checks score before writing to disk.",
    "FD-41": "Empty input must be a no-op, not an exception, or any quiet clip with no people would kill the run.",
    "FD-42": "max_faces_per_person=10 cap. New face is worse than the worst kept one, so we keep what we have.",
    "FD-43": "Quality-aware eviction is what makes the per-person folder a 'best 10' set rather than 'first 10'.",
    "FD-44": "When YOLO returns two boxes for one person in one frame, we keep only the cleaner detection.",

    # ---------- Face Tracking — OSNet ----------
    "FT-01": "OSNet must produce one 512-d vector per box; mismatched lengths silently misalign embeddings to people.",
    "FT-02": "Returning None at the slot keeps batch alignment intact even when one crop is too small to embed.",
    "FT-03": "Cheap early-exit: calling OSNet on a batch of zero-sized crops wastes GPU and produces noise embeddings.",
    "FT-04": "OSNet outputs are L2-normalised so cosine = inner product; without unit length, every distance threshold is wrong.",
    "FT-05": "GPU OOM or driver hiccups must not break a 5-hour run; the per-crop fallback keeps tracking alive.",
    "FT-06": "Mirror of FT-02 for the single-shot path; both branches must agree on the size rule.",
    "FT-07": "YOLO sometimes returns coordinates slightly off-frame; clipping prevents NumPy index errors deep inside OSNet.",
    # ---------- Face Tracking — ActiveTrack ----------
    "FT-08": "When a track is brand new, the reference IS the only embedding; no averaging tricks should kick in.",
    "FT-09": "A person's appearance shifts during a clip (lighting, angle); newer embeddings should outweigh older ones.",
    "FT-10": "The reference goes back into cosine math, so it must stay on the unit sphere even after weighted averaging.",
    "FT-11": "Without growth, the bank stays at size 1 forever and the recency weighting tested in FT-09 is meaningless.",
    "FT-12": "Bank cap protects RAM; evicting the oldest keeps the bank biased toward 'how this person looks right now'.",
    "FT-13": "Stops a one-frame mis-match (a stranger briefly walking through frame) from poisoning the track's identity.",
    "FT-14": "Cross-clip continuity needs the .npy path on disk; without bank_file, restore_state can't find the file.",
    "FT-15": "If the loaded bank produces a different reference, the same person across clips would mismatch their old ID.",
    "FT-16": "Crash recovery: a half-written state.json or deleted .npy must not kill the next clip's start-up.",
    "FT-17": "Catches subtle changes (dtype loss, dimension swap) that only show up after a serialise→deserialise round trip.",
    # ---------- Face Tracking — PendingTrack ----------
    "FT-18": "First detection counts as one frame; without this, warmup_frames=10 would actually need 11 to promote.",
    "FT-19": "Pending tracks use a simple mean (no recency weighting) because there isn't enough history to weight by.",
    "FT-20": "If the bank doesn't grow during pending, promotion's bank-seed (FT-36) has nothing to copy across.",
    "FT-21": "Pending → Active is the moment a person gets a real track ID; if it never fires, no one ever appears in the report.",
    "FT-22": "If promotion fired too early, every brief flicker would become a track and the report would explode with phantom people.",
    # ---------- Face Tracking — _box_distance ----------
    "FT-23": "Sanity floor: if two perfectly overlapping boxes don't have distance 0, every later distance comparison drifts.",
    "FT-24": "0.3 is the spatial gate; this confirms a clear teleport actually exceeds it (otherwise FT-29 wouldn't catch teleports).",
    "FT-25": "The same physical movement at different camera distances must give the same distance; otherwise zoom level breaks tracking.",
    "FT-26": "Locks the math down so a future refactor can't subtly change the centroid metric and break every threshold.",
    # ---------- Face Tracking — match_or_create ----------
    "FT-27": "The single most important tracker assertion: if this fails, every person becomes 50 different people in the report.",
    "FT-28": "When ReID says 'this isn't anyone we know', a new track must start instead of being force-matched to a stranger.",
    "FT-29": "Two people with similar clothes shouldn't swap IDs as they cross; the spatial gate enforces 'same person can't jump'.",
    "FT-30": "Between clips there's a real time gap; a strict gate would refuse all reconnections, so it widens with elapsed time.",
    "FT-31": "When confidence is high, fold the new appearance into the bank so the track stays current with reality.",
    "FT-32": "Don't poison a good track with a maybe-match; only confident matches earn a bank update.",
    "FT-33": "When 3 detections compete for 3 tracks, the assignment must be globally best — naive first-fit would mis-pair people.",
    "FT-34": "Without per-frame increments, pending tracks never reach the warmup threshold and nobody ever gets promoted.",
    "FT-35": "Mirror of FT-21 inside the full PersonTracker; the threshold must fire end-to-end, not just on the inner class.",
    "FT-36": "On promotion, the new active track must inherit the warmup-period history; otherwise it instantly 'forgets' the person.",
    "FT-37": "First time we see somebody, a pending track must be created instead of the detection being silently thrown away.",
    "FT-38": "An undersized crop returns None from OSNet; the tracker must skip that detection, not crash the frame.",
    "FT-39": "Frames with no people are common; the tracker must accept an empty input as a normal occurrence.",
    "FT-40": "When some people are temporarily occluded, only the visible ones get matched; the rest must remain until lost-timeout.",
    # ---------- Face Tracking — get_lost ----------
    "FT-41": "Lost-track GC: after max_lost_seconds the track must be reaped, freeing its ID for reuse and removing it from RAM.",
    "FT-42": "Don't reap healthy tracks; otherwise everyone would disappear from the report mid-day.",
    "FT-43": "Empty-state safety: get_lost on an empty tracker must not crash.",
    "FT-44": "state.json stores ISO timestamps; if parsing breaks, lost-detection silently never fires across reboots.",
    "FT-45": "Boundary: strict greater-than means a track exactly at 25.0 s gets one more chance instead of being killed.",
    # ---------- Face Tracking — serialize / restore / remove ----------
    "FT-46": "Mirror of FT-14 at the tracker level; without bank_file keys here, restore at the next clip can't find any banks.",
    "FT-47": "Cross-clip continuity end-to-end: same person across a clip boundary must re-attach to their old ID.",
    "FT-48": "If a bank file is deleted between clips, restore must skip it not crash; consistent with FT-16.",
    "FT-49": "Operator can blacklist a stuck or wrong track; without remove() that ID would persist forever.",
    "FT-50": "Idempotent remove: calling twice on the same ID shouldn't raise an exception.",
    # ---------- Face Tracking — LineCrossDetector ----------
    "FT-51": "Side classification underpins every crossing; if 'above' is wrongly classified, in/out counts invert.",
    "FT-52": "Twin of FT-51 on the negative side; both halves of the line must classify consistently.",
    "FT-53": "Points exactly on the line are ambiguous; returning 0 lets the update logic skip them safely.",
    "FT-54": "Vertical lines are a common store layout (door at the side); the math must not assume horizontal orientation.",
    "FT-55": "Core line-crossing event: without this, no crossing ever fires regardless of how people move.",
    "FT-56": "Same-side movement mustn't trigger a crossing; false positives would inflate the entry/exit count silently.",
    "FT-57": "We can't 'cross' on first sight (no prior side known); must store the baseline silently.",
    "FT-58": "Entry vs exit matters for the report; without direction, footfall is just 'movement count'.",
    "FT-59": "Without cleanup, every track ID ever seen sits in the crossing dict forever — slow and memory-hungry.",
    "FT-60": "Report needs entry/exit totals; this is the public read interface.",
    "FT-61": "After a clean state, get_state must return empty — otherwise yesterday's counts leak into today's report.",

    # ---------- Face Recognition — extract_embeddings ----------
    "FR-EX-01": "Baseline that FR ingest works at all; if zero embeddings come out of a normal folder, FR is a no-op.",
    "FR-EX-02": "no_face people are valid; ingest must return empty cleanly instead of erroring on an empty folder.",
    "FR-EX-03": "body_snapshot.jpg is a fallback only; including it in the regular face scan would skew embeddings with non-face data.",
    "FR-EX-04": "When all face crops are unusable, the body snapshot becomes the last identity hint; without it, the person becomes no_face.",
    "FR-EX-05": "Filesystems collect .DS_Store, Thumbs.db, half-written tmps; ingest must filter to .jpg only.",
    # ---------- Face Recognition — pick_diverse ----------
    "FR-PD-01": "Staff registry stores 'mean + most-distant'; without diversity, the registry over-fits one camera angle.",
    "FR-PD-02": "Cap of 4 'extras' protects FAISS index size; without a cap, a staff member with 100 photos would explode the index.",
    "FR-PD-03": "Locks the cosine-distance formula so a future refactor can't silently change the diversity metric.",
    "FR-PD-04": "Don't fail when a staff member only has 2 photos; just return what's available.",
    "FR-PD-05": "Cap=0 lets ops register staff with mean-only mode for tiny datasets where diversity is meaningless.",
    "FR-PD-06": "Re-running registration on the same photos must produce the same FAISS index; otherwise the index drifts every rebuild.",
    # ---------- Face Recognition — classify_with_faiss ----------
    "FR-SM-01": "The whole point of the staff registry: above-threshold matches must classify as staff.",
    "FR-SM-02": "False-staff classifications would tag every visitor as a known employee; the threshold has to actually reject strangers.",
    "FR-SM-03": "Documents whether the gate is >= or >; pinning the choice prevents silent drift on future tweaks.",
    "FR-SM-04": "no_face is its own category; without it, blind tracks would be force-classified as customers.",
    "FR-SM-05": "A person with many faces should be classified by their best match, giving staff a fighting chance even on bad clips.",
    "FR-SM-06": "Day-1 deployments have no staff registered yet; FR must still run and emit customers instead of crashing.",
    "FR-SM-07": "Defensive path: if FAISS row count mismatches registry length, classify must fail soft, not crash the report.",
    # ---------- Face Recognition — greedy_cluster ----------
    "FR-CC-01": "When the same person is captured under two track IDs (e.g. a brief occlusion), clustering should merge them.",
    "FR-CC-02": "Without this guarantee, clustering could collapse all customers into one giant blob.",
    "FR-CC-03": "Trivial case but locks the loop's empty-list and single-item handling.",
    "FR-CC-04": "Some clips have only staff; clustering must handle the empty input cleanly and return [].",
    "FR-CC-06": "Locks the cluster gate's inclusivity (>=); flipping it changes how many clusters appear in every report.",
    "FR-CC-07": "no_face people have no embedding to compare; clustering would crash on a None vector without this filter.",
    "FR-CC-08": "Sanity case — collapsing duplicates into one cluster is exactly what clustering is for.",
}


# ----------------------------------------------------------------------------
# Colour palette
# ----------------------------------------------------------------------------
NAVY = RGBColor(0x0E, 0x2A, 0x47)
ACCENT = RGBColor(0x2F, 0x54, 0x96)         # FD blue
ACCENT_GREEN = RGBColor(0x2E, 0x7D, 0x32)
ACCENT_AMBER = RGBColor(0xC8, 0x7F, 0x0A)
ACCENT_RED = RGBColor(0xC0, 0x39, 0x2B)
ACCENT_PURPLE = RGBColor(0x6B, 0x2F, 0xA0)  # FT
ACCENT_TEAL = RGBColor(0x0F, 0x8B, 0x8D)    # FR
LIGHT = RGBColor(0xF4, 0xF6, 0xFA)
GREY = RGBColor(0x55, 0x5C, 0x66)
BORDER = RGBColor(0xCF, 0xD6, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ----------------------------------------------------------------------------
# Slide setup (16:9)
# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def add_rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, *, size=14, bold=False, color=NAVY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, *, size=14, color=NAVY, bullet="•"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = f"{bullet}  {item}"
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_slide_header(slide, title, subtitle=None, accent=ACCENT):
    add_rect(slide, 0, 0, SW, Inches(0.9), accent)
    add_rect(slide, 0, Inches(0.9), SW, Inches(0.04), NAVY)
    add_text(slide, Inches(0.5), Inches(0.18), SW - Inches(1), Inches(0.55),
             title, size=26, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.6), SW - Inches(1), Inches(0.3),
                 subtitle, size=12, color=RGBColor(0xDD, 0xE6, 0xF2))
    # footer band
    add_rect(slide, 0, SH - Inches(0.32), SW, Inches(0.32), LIGHT)
    add_text(slide, Inches(0.5), SH - Inches(0.3), Inches(8), Inches(0.28),
             "FR_Thor — Offline Face Recognition Pipeline  |  Unit Testing",
             size=10, color=GREY)
    add_text(slide, SW - Inches(2.5), SH - Inches(0.3), Inches(2), Inches(0.28),
             f"Run: {STATE['run_date']}", size=10, color=GREY, align=PP_ALIGN.RIGHT)


def add_table(slide, x, y, w, h, headers, rows, *,
              header_fill=NAVY, header_color=WHITE,
              row_fill=WHITE, alt_fill=LIGHT, body_color=NAVY,
              header_size=12, body_size=11, col_widths=None,
              verdict_col=None):
    n_cols = len(headers)
    n_rows = len(rows) + 1
    table_shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = table_shape.table
    if col_widths:
        total = sum(col_widths)
        for i, ratio in enumerate(col_widths):
            table.columns[i].width = int(w * ratio / total)
    for j, head in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        cell.text = ""
        tf = cell.text_frame
        tf.margin_left = Emu(60000)
        tf.margin_right = Emu(60000)
        tf.margin_top = Emu(40000)
        tf.margin_bottom = Emu(40000)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = head
        run.font.size = Pt(header_size)
        run.font.bold = True
        run.font.color.rgb = header_color
        run.font.name = "Calibri"
    for i, row in enumerate(rows, start=1):
        fill = row_fill if i % 2 == 1 else alt_fill
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.text = ""
            tf = cell.text_frame
            tf.margin_left = Emu(60000)
            tf.margin_right = Emu(60000)
            tf.margin_top = Emu(30000)
            tf.margin_bottom = Emu(30000)
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(body_size)
            colour = body_color
            bold = False
            if verdict_col is not None and j == verdict_col:
                v = str(val).upper()
                if v == "PASS":
                    colour = ACCENT_GREEN
                    bold = True
                elif v == "FAIL":
                    colour = ACCENT_RED
                    bold = True
                elif v in ("SKIP", "ERROR", "NOT RUN"):
                    colour = ACCENT_AMBER
                    bold = True
            run.font.bold = bold
            run.font.color.rgb = colour
            run.font.name = "Calibri"
    return table


def add_metric_card(slide, x, y, w, h, label, value, accent):
    add_rect(slide, x, y, w, h, WHITE, line=BORDER)
    add_rect(slide, x, y, Inches(0.12), h, accent)
    add_text(slide, x + Inches(0.3), y + Inches(0.15), w - Inches(0.4), Inches(0.3),
             label, size=12, color=GREY, bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(0.45), w - Inches(0.4), h - Inches(0.5),
             str(value), size=36, color=NAVY, bold=True)


# ----------------------------------------------------------------------------
# Cover / structural slides
# ----------------------------------------------------------------------------
def slide_title():
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, SW, SH, NAVY)
    add_rect(s, 0, Inches(2.4), SW, Inches(0.06), ACCENT)
    add_rect(s, 0, Inches(4.8), SW, Inches(0.06), ACCENT_TEAL)
    add_text(s, Inches(0.7), Inches(2.7), Inches(12), Inches(1.2),
             "FR_Thor — Unit Testing", size=54, bold=True, color=WHITE)
    add_text(s, Inches(0.7), Inches(3.85), Inches(12), Inches(0.5),
             "Offline Face Recognition Pipeline · Validation Walk-through",
             size=22, color=RGBColor(0xCD, 0xDA, 0xE9))
    chip_y = Inches(5.3)
    chips = [
        ("130 Test Cases", ACCENT),
        ("3 Features Covered", ACCENT_TEAL),
        ("100% PASS", ACCENT_GREEN),
        ("New-Joinee Friendly", ACCENT_PURPLE),
    ]
    cx = Inches(0.7)
    for label, col in chips:
        w = Inches(2.6)
        add_rect(s, cx, chip_y, w, Inches(0.55), col)
        add_text(s, cx, chip_y, w, Inches(0.55), label,
                 size=14, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cx += w + Inches(0.2)
    add_text(s, Inches(0.7), SH - Inches(0.7), Inches(12), Inches(0.4),
             f"Pipeline v1.7.0  ·  Test Run {STATE['run_ts']}  ·  Real-Footage Integration",
             size=12, color=RGBColor(0xB6, 0xC4, 0xD8))


def slide_agenda():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Agenda", "What we'll walk through together")
    items = [
        ("1", "Project & Pipeline Overview", "What FR_Thor does, end-to-end"),
        ("2", "Why We Test This Pipeline", "Risks the suite is designed to catch"),
        ("3", "Testing Approach", "Real-footage integration, no mocks"),
        ("4", "How to Read These Slides", "ID · what · why it matters · verdict"),
        ("5", "Face Detection — 44 cases", "Filter garbage out before FR sees it"),
        ("6", "Face Tracking — 61 cases", "Give every visitor one stable ID"),
        ("7", "Face Recognition — 25 cases", "Tell staff and customers apart"),
        ("8", "Tooling, Environment, Results", "How to run it, what we got, what's next"),
    ]
    y = Inches(1.2)
    for n, title, desc in items:
        add_rect(s, Inches(0.7), y, Inches(0.6), Inches(0.6), ACCENT)
        add_text(s, Inches(0.7), y, Inches(0.6), Inches(0.6),
                 n, size=22, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(1.5), y + Inches(0.05), Inches(11), Inches(0.32),
                 title, size=16, bold=True, color=NAVY)
        add_text(s, Inches(1.5), y + Inches(0.32), Inches(11), Inches(0.3),
                 desc, size=12, color=GREY)
        y += Inches(0.65)


def slide_project_overview():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Project Overview", "FR_Thor at a glance")
    add_text(s, Inches(0.7), Inches(1.2), Inches(6), Inches(0.4),
             "What is FR_Thor?", size=18, bold=True, color=NAVY)
    add_text(s, Inches(0.7), Inches(1.65), Inches(6), Inches(2.4),
             ("An offline face-recognition pipeline for pharmacy footfall "
              "analysis. It pulls recorded CCTV clips from a remote Jetson, "
              "tracks people across clips, captures their faces, identifies "
              "staff against a registry, clusters unknown visitors, and "
              "produces a multi-sheet Excel report.\n\n"
              "All inference runs locally — no cloud APIs, no PII leaving "
              "the store."),
             size=13, color=NAVY)
    add_text(s, Inches(7.2), Inches(1.2), Inches(5.5), Inches(0.4),
             "Core Models & Stack", size=18, bold=True, color=NAVY)
    rows = [
        ["YOLO11l", "Person detection (conf 0.45)"],
        ["YOLOv11s-face", "Face detection (conf 0.30)"],
        ["OSNet (boxmot)", "512-d ReID embeddings"],
        ["InsightFace antelopev2", "ArcFace 512-d face embeddings"],
        ["FAISS IndexFlatIP", "Cosine similarity search"],
        ["state.json", "Per-date crash-safe persistence"],
    ]
    add_table(s, Inches(7.2), Inches(1.65), Inches(5.5), Inches(2.6),
              ["Component", "Role"], rows,
              col_widths=[2, 3.5], header_size=11, body_size=11)
    add_metric_card(s, Inches(0.7), Inches(4.7), Inches(2.9), Inches(1.4),
                    "Pipeline Version", "v1.7.0", ACCENT)
    add_metric_card(s, Inches(3.8), Inches(4.7), Inches(2.9), Inches(1.4),
                    "Pipeline Phases", "4", ACCENT_TEAL)
    add_metric_card(s, Inches(6.9), Inches(4.7), Inches(2.9), Inches(1.4),
                    "Persistence", "JSON only", ACCENT_PURPLE)
    add_metric_card(s, Inches(10.0), Inches(4.7), Inches(2.7), Inches(1.4),
                    "Inference", "100% local", ACCENT_GREEN)


def slide_pipeline_arch():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Pipeline Architecture", "Four phases, fully offline")
    phases = [
        ("Tracking",
         "YOLO person detection → OSNet ReID → PersonTracker (3-phase greedy match)",
         ACCENT),
        ("Face Capture",
         "YOLO face detection → score_face quality → async save face_*_qNN.jpg",
         ACCENT_TEAL),
        ("Face Recognition",
         "3-filter (quality / skin / keypoints) → InsightFace embed → FAISS staff match → cluster",
         ACCENT_PURPLE),
        ("Reporting",
         "6-sheet Excel report with thumbnails per person, written to data/<date>/report/",
         ACCENT_GREEN),
    ]
    y = Inches(1.3)
    for i, (name, desc, col) in enumerate(phases):
        add_rect(s, Inches(0.7), y, Inches(2.6), Inches(1.0), col)
        add_text(s, Inches(0.7), y, Inches(2.6), Inches(1.0),
                 f"Phase {i+1}\n{name}", size=16, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, Inches(3.4), y, Inches(9.3), Inches(1.0), WHITE, line=BORDER)
        add_text(s, Inches(3.6), y + Inches(0.05), Inches(9.0), Inches(0.95),
                 desc, size=13, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(phases) - 1:
            arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                                       Inches(1.85), y + Inches(1.05),
                                       Inches(0.3), Inches(0.18))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GREY
            arrow.line.fill.background()
        y += Inches(1.22)


def slide_why_test():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Why Unit-Test This Pipeline?",
                     "Each stage has its own way of going silently wrong")
    risks = [
        ("Detection drift",
         "A new YOLO weight or a new conf value changes which crops are kept. "
         "Without tests, the only signal is 'reports look weird'."),
        ("Tracking ID swaps",
         "A small math change in OSNet or the spatial gate can make two "
         "different people share an ID — every footfall metric becomes wrong."),
        ("FR mis-classification",
         "A tweak to the FAISS threshold can flip every customer to staff or "
         "vice-versa. The Excel report still 'looks fine' but the data lies."),
        ("Silent data loss",
         "An off-by-one in the q-score gate, the skin filter or the cooldown "
         "can drop 90% of valid faces with no exception or warning."),
        ("Cross-clip continuity",
         "If serialise/restore or get_lost stops working, every clip starts "
         "fresh and one customer becomes hundreds of one-clip ghosts."),
    ]
    y = Inches(1.2)
    for title, desc in risks:
        add_rect(s, Inches(0.6), y, Inches(0.18), Inches(0.95), ACCENT_RED)
        add_rect(s, Inches(0.78), y, Inches(12.0), Inches(0.95), WHITE, line=BORDER)
        add_text(s, Inches(0.95), y + Inches(0.07), Inches(11.5), Inches(0.32),
                 title, size=14, bold=True, color=NAVY)
        add_text(s, Inches(0.95), y + Inches(0.4), Inches(11.5), Inches(0.55),
                 desc, size=12, color=NAVY)
        y += Inches(1.05)


def slide_testing_approach():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Testing Approach", "Real footage over mocks, every time")
    col_w = Inches(6.0)
    gap = Inches(0.3)
    x1 = Inches(0.7)
    x2 = x1 + col_w + gap
    add_rect(s, x1, Inches(1.2), col_w, Inches(2.7), WHITE, line=BORDER)
    add_rect(s, x1, Inches(1.2), col_w, Inches(0.55), ACCENT)
    add_text(s, x1, Inches(1.2), col_w, Inches(0.55),
             "  Real-Footage Integration", size=16, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, x1 + Inches(0.3), Inches(1.9), col_w - Inches(0.4), Inches(2),
                ["Real CCTV clips from data/2026-04-29/",
                 "Real models — YOLO, OSNet, InsightFace",
                 "Tests assert behaviour on actual frames",
                 "No mocking of detection or embedding paths",
                 "Catches model-version & data-shape regressions"],
                size=13)
    add_rect(s, x2, Inches(1.2), col_w, Inches(2.7), WHITE, line=BORDER)
    add_rect(s, x2, Inches(1.2), col_w, Inches(0.55), ACCENT_TEAL)
    add_text(s, x2, Inches(1.2), col_w, Inches(0.55),
             "  Targeted, Per-Function Cases", size=16, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, x2 + Inches(0.3), Inches(1.9), col_w - Inches(0.4), Inches(2),
                ["One script per test ID (test_fdNN, test_ftNN…)",
                 "Each script writes its own results.json",
                 "Master runner aggregates → state JSON + Excel",
                 "Easy to re-run a single failing case",
                 "Replayable from saved state JSON"],
                size=13)
    add_rect(s, Inches(0.7), Inches(4.2), Inches(12.0), Inches(2.5), LIGHT, line=BORDER)
    add_text(s, Inches(1.0), Inches(4.35), Inches(11.5), Inches(0.4),
             "Source files are read-only", size=15, bold=True, color=NAVY)
    add_text(s, Inches(1.0), Inches(4.75), Inches(11.5), Inches(1.85),
             ("The tests live alongside the code but never touch it. "
              "pipeline.py, fr_report.py, rerun_fr.py, register_staff.py and "
              "filter_faces.py are treated as the system under test — tests "
              "verify behaviour from the outside.\n\n"
              "Test scripts, plans and run reports sit under Unit_Test_Data/ "
              "and Unit_Test_Test_Cases/. Source code never changes to make a "
              "test pass."),
             size=12, color=NAVY)


def slide_how_to_read():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "How to Read the Test-Case Slides",
                     "Same shape every time, so you can skim or deep-read")
    add_text(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.4),
             "Every feature slide that follows uses this layout:",
             size=14, color=NAVY)
    # Sample table mimicking the real one
    rows = [
        ["FD-01", "Low-confidence face detection should be rejected outright",
         "Without this gate, every wall, shadow or poster YOLO mis-flags would be saved as a face.",
         "PASS"],
    ]
    add_table(s, Inches(0.7), Inches(1.7), Inches(12.0), Inches(1.2),
              ["ID", "What it tests", "Why it matters", "Verdict"], rows,
              col_widths=[1.0, 4.5, 5.7, 0.8],
              header_size=12, body_size=11, verdict_col=3)
    # Legend cards
    items = [
        ("ID", "Stable test identifier — also the name of the script "
               "in Unit_Test_Data/Face_*/scripts/.", ACCENT),
        ("What it tests", "Plain-English description of the input and "
                          "the assertion. No code lingo.", ACCENT_TEAL),
        ("Why it matters", "What real-world bug or regression this case "
                           "catches. Read this first when reviewing results.", ACCENT_PURPLE),
        ("Verdict", "PASS · FAIL · SKIP · ERROR. Today every case is PASS "
                    "(green); failures would be red, skips amber.", ACCENT_GREEN),
    ]
    y = Inches(3.1)
    for label, desc, col in items:
        add_rect(s, Inches(0.7), y, Inches(2.4), Inches(0.85), col)
        add_text(s, Inches(0.7), y, Inches(2.4), Inches(0.85),
                 label, size=15, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_rect(s, Inches(3.2), y, Inches(9.5), Inches(0.85), WHITE, line=BORDER)
        add_text(s, Inches(3.4), y + Inches(0.07), Inches(9.2), Inches(0.75),
                 desc, size=12, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.95)


def slide_suite_overview():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Test Suite Overview",
                     "Three feature areas, 130 cases total")
    fd = STATE["features"]["face_detection"]["counts"]
    ft = STATE["features"]["face_tracking"]["counts"]
    fr = STATE["features"]["face_recognition"]["counts"]
    total = fd["PASS"] + ft["PASS"] + fr["PASS"]
    add_metric_card(s, Inches(0.7), Inches(1.2), Inches(2.95), Inches(1.4),
                    "Total Tests", total, NAVY)
    add_metric_card(s, Inches(3.75), Inches(1.2), Inches(2.95), Inches(1.4),
                    "Face Detection", fd["PASS"], ACCENT)
    add_metric_card(s, Inches(6.8), Inches(1.2), Inches(2.95), Inches(1.4),
                    "Face Tracking", ft["PASS"], ACCENT_PURPLE)
    add_metric_card(s, Inches(9.85), Inches(1.2), Inches(2.85), Inches(1.4),
                    "Face Recognition", fr["PASS"], ACCENT_TEAL)
    add_text(s, Inches(0.7), Inches(2.85), Inches(12), Inches(0.4),
             "Run Outcome", size=16, bold=True, color=NAVY)
    bar_y = Inches(3.3)
    bar_h = Inches(0.55)
    add_rect(s, Inches(0.7), bar_y, Inches(12), bar_h, ACCENT_GREEN)
    add_text(s, Inches(0.7), bar_y, Inches(12), bar_h,
             f"  PASS {total} / {total}   ·   FAIL 0   ·   SKIP 0   ·   ERROR 0",
             size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    rows = [
        ["Face Detection", "44", "score_face, _check_skin_tone, _count_kps_in_bbox, _batch_capture_faces"],
        ["Face Tracking",  "61", "OSNetExtractor, ActiveTrack, PendingTrack, PersonTracker, LineCrossDetector"],
        ["Face Recognition", "25", "extract_embeddings, _pick_diverse, classify_with_faiss, greedy_cluster"],
    ]
    add_text(s, Inches(0.7), Inches(4.15), Inches(12), Inches(0.4),
             "Coverage by Feature", size=16, bold=True, color=NAVY)
    add_table(s, Inches(0.7), Inches(4.6), Inches(12.0), Inches(2.0),
              ["Feature", "Cases", "Functions Under Test"], rows,
              col_widths=[2.5, 1.0, 7.5], body_size=12)


# ----------------------------------------------------------------------------
# Generic builders for feature intros and per-function case slides
# ----------------------------------------------------------------------------
def feature_intro_slide(title, subtitle, accent, where_it_sits, failure_modes,
                        why_it_matters):
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, title, subtitle, accent=accent)
    # Where it sits
    add_text(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.4),
             "Where it sits in the pipeline", size=15, bold=True, color=NAVY)
    add_rect(s, Inches(0.7), Inches(1.6), Inches(12.0), Inches(0.85),
             LIGHT, line=BORDER)
    add_text(s, Inches(0.95), Inches(1.7), Inches(11.5), Inches(0.7),
             where_it_sits, size=12, color=NAVY,
             anchor=MSO_ANCHOR.MIDDLE)
    # Failure modes
    add_text(s, Inches(0.7), Inches(2.65), Inches(12), Inches(0.4),
             "What goes wrong if this feature breaks", size=15, bold=True, color=NAVY)
    add_bullets(s, Inches(0.95), Inches(3.05), Inches(12), Inches(2.0),
                failure_modes, size=12)
    # Why this set of tests
    add_text(s, Inches(0.7), Inches(5.4), Inches(12), Inches(0.4),
             "Why these tests, in this shape", size=15, bold=True, color=NAVY)
    add_text(s, Inches(0.95), Inches(5.8), Inches(12.0), Inches(1.3),
             why_it_matters, size=12, color=NAVY)


def feature_functions_slide(title, subtitle, accent, function_rows):
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, title, subtitle, accent=accent)
    add_text(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.4),
             "Functions under test in this feature", size=15, bold=True, color=NAVY)
    add_table(s, Inches(0.7), Inches(1.65), Inches(12.0), Inches(5.2),
              ["Function", "Cases", "Role in the pipeline"],
              function_rows,
              col_widths=[3.2, 1.0, 6.6], header_size=12, body_size=11)


def function_cases_slide(title, accent, intro_line, test_ids, *, slide_label=""):
    """One slide showing the cases for one function (or one chunk of a function)."""
    s = prs.slides.add_slide(BLANK)
    n = len(test_ids)
    case_range = f"{test_ids[0]} – {test_ids[-1]}" if n > 1 else test_ids[0]
    suffix = f"  ·  {slide_label}" if slide_label else ""
    add_slide_header(s, title, f"{n} case{'s' if n != 1 else ''}  ·  {case_range}{suffix}",
                     accent=accent)
    # "What this function does" intro card
    card_y = Inches(1.05)
    add_rect(s, Inches(0.5), card_y, Inches(12.3), Inches(0.95),
             LIGHT, line=BORDER)
    add_text(s, Inches(0.75), card_y + Inches(0.07), Inches(12), Inches(0.35),
             "What this function does", size=13, bold=True, color=NAVY)
    add_text(s, Inches(0.75), card_y + Inches(0.42), Inches(12), Inches(0.5),
             intro_line, size=12, color=NAVY)
    # Cases table
    rows = []
    for tid in test_ids:
        t = find_test(tid)
        rows.append([
            tid,
            t["description"],
            RATIONALE.get(tid, "—"),
            t["verdict"],
        ])
    table_h = Inches(0.45) + Inches(0.65) * n   # rough auto-fit
    if table_h > Inches(4.6):
        table_h = Inches(4.6)
    add_table(s, Inches(0.5), Inches(2.2), Inches(12.3), table_h,
              ["ID", "What it tests", "Why it matters", "Verdict"], rows,
              col_widths=[1.0, 4.5, 5.7, 0.8],
              header_size=12, body_size=10, verdict_col=3)


# ----------------------------------------------------------------------------
# Face Detection slides
# ----------------------------------------------------------------------------
def fd_intro():
    feature_intro_slide(
        "Face Detection — Why It Matters",
        "44 cases  ·  the first quality gate of the pipeline",
        ACCENT,
        where_it_sits=(
            "Phase 2 of the pipeline. After YOLO finds a face inside a person "
            "crop, FD decides whether the crop is good enough to save. "
            "Everything Face Recognition sees later was approved here first."),
        failure_modes=[
            "Garbage-in / garbage-out: a permissive FD lets walls and posters "
            "into FR, polluting clusters and Excel reports.",
            "Over-strict FD drops valid faces silently — visitors disappear "
            "from the report with no error message.",
            "Quality-aware eviction is what keeps each person's folder a "
            "'best 10 faces' set; if it breaks, only the first 10 frames win.",
            "Strategy fallbacks (upscale / CLAHE / pad) are how we recover "
            "small or low-contrast crops; if any rung breaks, FR yield drops.",
        ],
        why_it_matters=(
            "Face Detection is split across many small helpers (score_face, "
            "skin-tone, keypoint containment, capture, embed-with-strategies). "
            "We test each helper on real frames so that swapping a YOLO weight "
            "or tuning a conf value can never silently change which faces "
            "survive into FR.")
    )


def fd_functions():
    feature_functions_slide(
        "Face Detection — Functions Under Test",
        "Each row maps a real function to the cases that exercise it",
        ACCENT,
        function_rows=[
            ["score_face", "11", "Frame-level quality 0–100+ (blur, size, brightness, pose)"],
            ["check_quality", "4", "Reads the q-tag baked into the face filename"],
            ["check_skin_tone / _check_skin_tone", "5",
             "HSV skin-pixel ratio over the centre 60% of the crop"],
            ["check_keypoints", "3", "Counts InsightFace landmarks inside the YOLO box"],
            ["_count_kps_in_bbox", "4", "Pure-math helper that powers the keypoint check"],
            ["_enhance_contrast", "2", "CLAHE rescue used inside the 4-strategy embed ladder"],
            ["_detect_face", "3", "Largest-face pick + L2 normalisation; returns None on miss"],
            ["_extract_with_strategies", "5", "Direct → upscale → CLAHE → padded fallback chain"],
            ["_batch_capture_faces", "7", "Async save + cooldown + max-faces eviction"],
        ])


def fd_score_face_1():
    function_cases_slide(
        "score_face — Confidence & size gates",
        ACCENT,
        ("Returns an integer quality score for a face crop. Inputs that fail "
         "the conf or size floor get score 0; everything else is graded by "
         "blur, brightness and pose sub-scores."),
        ["FD-01", "FD-02", "FD-03", "FD-04"],
        slide_label="1 of 3")


def fd_score_face_2():
    function_cases_slide(
        "score_face — Overall score & sub-scores",
        ACCENT,
        ("The overall quality score is the sum of blur, size, brightness and "
         "pose. These cases pin the upper end of the range and the behaviour "
         "of the blur and pose components on real crops."),
        ["FD-05", "FD-06", "FD-07", "FD-08"],
        slide_label="2 of 3")


def fd_score_face_3():
    function_cases_slide(
        "score_face — Keypoint fallbacks & brightness",
        ACCENT,
        ("These cases prove score_face stays sane when YOLO's pose head is "
         "missing or unreliable, and that overexposed frames don't sneak past "
         "the brightness component."),
        ["FD-09", "FD-10", "FD-11"],
        slide_label="3 of 3")


def fd_check_quality():
    function_cases_slide(
        "check_quality — Filename q-tag gate",
        ACCENT,
        ("Reads the integer score from filenames like face_001234_q72.jpg. "
         "If the score is below 25 the face is rejected; missing tag is "
         "treated as 'pass'."),
        ["FD-12", "FD-13", "FD-14", "FD-15"])


def fd_skin():
    function_cases_slide(
        "check_skin_tone / _check_skin_tone — HSV skin filter",
        ACCENT,
        ("Cheap second filter: convert the centre 60% of the crop to HSV and "
         "demand at least 20% skin-coloured pixels. Mostly drops walls, "
         "jackets and signs that YOLO mis-classifies as faces."),
        ["FD-16", "FD-17", "FD-18", "FD-22", "FD-23"])


def fd_keypoints_outer():
    function_cases_slide(
        "check_keypoints — Landmark containment (3rd filter)",
        ACCENT,
        ("Runs InsightFace on the crop and demands that at least 3 of its 5 "
         "landmarks fall inside the YOLO bounding box. If they don't, the "
         "crop is shaped like something else (hat, scarf, sign)."),
        ["FD-19", "FD-20", "FD-21"])


def fd_kps_in_bbox():
    function_cases_slide(
        "_count_kps_in_bbox — Pure-math helper",
        ACCENT,
        ("Given a list of landmark coordinates and a bounding box, returns "
         "the count of points that fall inside the box (with a small margin). "
         "Used by check_keypoints."),
        ["FD-24", "FD-25", "FD-26", "FD-27"])


def fd_enhance_contrast():
    function_cases_slide(
        "_enhance_contrast — CLAHE rescue",
        ACCENT,
        ("Applies CLAHE (contrast-limited adaptive histogram equalisation) "
         "to a face crop so InsightFace can embed dark or low-contrast "
         "frames. Used as the third rung of the FR strategy ladder."),
        ["FD-28", "FD-29"])


def fd_detect_face():
    function_cases_slide(
        "_detect_face — Largest face + unit normalisation",
        ACCENT,
        ("Wraps InsightFace.get(): if multiple faces are found in the crop, "
         "pick the largest; L2-normalise the embedding so cosine similarity "
         "works; return None when no face is detected."),
        ["FD-30", "FD-31", "FD-32"])


def fd_strategies():
    function_cases_slide(
        "_extract_with_strategies — 4-strategy embedding ladder",
        ACCENT,
        ("Tries to embed a crop using four strategies in order: direct, "
         "upscaled, CLAHE-enhanced, padded. The first one that produces an "
         "embedding wins; returns (None, None) if all four fail."),
        ["FD-33", "FD-34", "FD-35", "FD-36", "FD-37"])


def fd_batch_capture_1():
    function_cases_slide(
        "_batch_capture_faces — Cooldown & basic gates",
        ACCENT,
        ("The capture stage that actually writes face_*_qNN.jpg files. "
         "Enforces a per-track cooldown and re-checks score before saving."),
        ["FD-38", "FD-39", "FD-40", "FD-41"],
        slide_label="1 of 2")


def fd_batch_capture_2():
    function_cases_slide(
        "_batch_capture_faces — Eviction & duplicate handling",
        ACCENT,
        ("When a track already has max_faces_per_person=10 stored, capture "
         "either evicts the worst face if the new one is better, or skips. "
         "Multi-detection frames keep only the highest-confidence box."),
        ["FD-42", "FD-43", "FD-44"],
        slide_label="2 of 2")


# ----------------------------------------------------------------------------
# Face Tracking slides
# ----------------------------------------------------------------------------
def ft_intro():
    feature_intro_slide(
        "Face Tracking — Why It Matters",
        "61 cases  ·  give every visitor one stable ID across frames and clips",
        ACCENT_PURPLE,
        where_it_sits=(
            "Phase 1 of the pipeline. YOLO sees a person every frame; the "
            "tracker decides whether 'that person' is the same one as in the "
            "previous frame, or a new visitor. Every later stage (FR, the "
            "Excel report, the entry/exit count) is keyed on that decision."),
        failure_modes=[
            "ID swaps: two people get the same track ID, so their faces get "
            "merged into one customer record.",
            "ID fragmentation: one person gets dozens of IDs across a clip, "
            "so the report shows phantom visitors and inflated footfall.",
            "Cross-clip drift: serialise/restore breaks, so every clip starts "
            "fresh and the same regular customer becomes 50 different ones.",
            "Lost-track leak: get_lost stops reaping idle tracks, RAM grows "
            "until the pipeline OOMs mid-run.",
        ],
        why_it_matters=(
            "Tracking has three layered components — OSNet for embeddings, "
            "ActiveTrack/PendingTrack for per-person state, PersonTracker "
            "for the matching loop — plus the LineCrossDetector for "
            "footfall counting. We test each layer on its own contract so a "
            "change in one layer can't silently break the others.")
    )


def ft_functions():
    feature_functions_slide(
        "Face Tracking — Functions Under Test",
        "Each row maps a real function to the cases that exercise it",
        ACCENT_PURPLE,
        function_rows=[
            ["OSNetExtractor.extract_batch / extract", "7",
             "Produce 512-d ReID embeddings per person crop"],
            ["ActiveTrack.ref / update_embedding", "6",
             "Recency-weighted reference + bank update gate"],
            ["ActiveTrack.serialize / deserialize", "4",
             ".npy round-trip for cross-clip continuity"],
            ["PendingTrack (init / ref / add / promotion)", "5",
             "Two-tier tracking: warmup before getting an ID"],
            ["PersonTracker._box_distance", "4",
             "Scale-invariant centroid distance for the spatial gate"],
            ["PersonTracker.match_or_create", "14",
             "3-phase greedy match (active → pending → new)"],
            ["PersonTracker.get_lost", "5",
             "Reap tracks past max_lost_seconds"],
            ["PersonTracker.serialize_state / restore_state / remove", "5",
             "End-of-clip persistence and cleanup"],
            ["LineCrossDetector.get_side / update / cleanup / get_state", "11",
             "Footfall counting on virtual entry/exit lines"],
        ])


def ft_osnet_batch():
    function_cases_slide(
        "OSNetExtractor.extract_batch — Batched ReID embedder",
        ACCENT_PURPLE,
        ("Takes a list of person crops and returns a list of 512-d unit "
         "vectors (one per box). Crops below the size floor get None at "
         "their slot so the list stays aligned with the input."),
        ["FT-01", "FT-02", "FT-03", "FT-04", "FT-05"])


def ft_osnet_single():
    function_cases_slide(
        "OSNetExtractor.extract — Single-shot embedder",
        ACCENT_PURPLE,
        ("Embeds a single crop. Same size rule and out-of-bounds handling "
         "as extract_batch — these tests pin those edges."),
        ["FT-06", "FT-07"])


def ft_active_ref():
    function_cases_slide(
        "ActiveTrack.ref — Recency-weighted reference",
        ACCENT_PURPLE,
        ("Returns the unit-normalised reference embedding for a track. "
         "Uses a recency-weighted average so the reference reflects the "
         "person's current appearance, not their first-frame appearance."),
        ["FT-08", "FT-09", "FT-10"])


def ft_active_update():
    function_cases_slide(
        "ActiveTrack.update_embedding — Bank update gate",
        ACCENT_PURPLE,
        ("Adds an embedding to the bank (capped at bank_size=10, oldest "
         "evicted) but only if it's close enough to the current reference "
         "to be plausibly the same person."),
        ["FT-11", "FT-12", "FT-13"])


def ft_active_serialise():
    function_cases_slide(
        "ActiveTrack.serialize / deserialize — Disk round-trip",
        ACCENT_PURPLE,
        ("At end-of-clip every active track's bank is written to "
         "tracks/track_XXXX_bank.npy and a 'bank_file' key is added to "
         "state.json. At the start of the next clip the bank is restored."),
        ["FT-14", "FT-15", "FT-16", "FT-17"])


def ft_pending():
    function_cases_slide(
        "PendingTrack — Warmup before promotion",
        ACCENT_PURPLE,
        ("New detections start as pending tracks (no ID yet) and only get "
         "promoted to ActiveTrack after warmup_frames=10 consecutive "
         "matches. This filter keeps brief flickers from becoming visitors."),
        ["FT-18", "FT-19", "FT-20", "FT-21", "FT-22"])


def ft_box_distance():
    function_cases_slide(
        "PersonTracker._box_distance — Scale-invariant centroid distance",
        ACCENT_PURPLE,
        ("Computes a normalised distance between two bounding-box centroids. "
         "The result is in 0..1 across the frame and is used as a spatial "
         "gate: even if ReID says 'same person', a too-large jump blocks "
         "the match."),
        ["FT-23", "FT-24", "FT-25", "FT-26"])


def ft_match_create_1():
    function_cases_slide(
        "PersonTracker.match_or_create — Match against active tracks",
        ACCENT_PURPLE,
        ("The main matching loop. Phase 1 tries to greedy-match new "
         "detections to existing active tracks using ReID + the spatial "
         "gate. These cases pin the basic match path and the gate's role."),
        ["FT-27", "FT-28", "FT-29", "FT-30", "FT-31"],
        slide_label="1 of 3")


def ft_match_create_2():
    function_cases_slide(
        "PersonTracker.match_or_create — Bank updates & pending lifecycle",
        ACCENT_PURPLE,
        ("Phase 2: matches that survive feed back into the bank (or don't, "
         "depending on confidence). Unmatched detections become pending; "
         "pending tracks promote once they reach warmup_frames."),
        ["FT-32", "FT-33", "FT-34", "FT-35", "FT-36"],
        slide_label="2 of 3")


def ft_match_create_3():
    function_cases_slide(
        "PersonTracker.match_or_create — New tracks & edge cases",
        ACCENT_PURPLE,
        ("Phase 3: detections with no match anywhere become brand-new "
         "pending tracks. These cases also lock down the empty-input and "
         "size-mismatch behaviour so the loop never crashes mid-clip."),
        ["FT-37", "FT-38", "FT-39", "FT-40"],
        slide_label="3 of 3")


def ft_get_lost():
    function_cases_slide(
        "PersonTracker.get_lost — Reap idle tracks",
        ACCENT_PURPLE,
        ("Returns the IDs of tracks whose last_seen is older than "
         "max_lost_seconds=25. These tracks are then finalised and their "
         "IDs freed. Without this, tracks would accumulate forever."),
        ["FT-41", "FT-42", "FT-43", "FT-44", "FT-45"])


def ft_serialize_remove():
    function_cases_slide(
        "PersonTracker — serialize_state / restore_state / remove",
        ACCENT_PURPLE,
        ("End-of-clip: serialize_state writes every active track's bank "
         "file path into state.json. Start-of-next-clip: restore_state "
         "loads them back. remove() lets ops delete a stuck track by ID."),
        ["FT-46", "FT-47", "FT-48", "FT-49", "FT-50"])


def ft_lcd_side():
    function_cases_slide(
        "LineCrossDetector.get_side — Which side of the line is this point?",
        ACCENT_PURPLE,
        ("Returns +1, -1, or 0 depending on which side of a virtual line "
         "a point sits. Foundation for entry/exit counting."),
        ["FT-51", "FT-52", "FT-53", "FT-54"])


def ft_lcd_update():
    function_cases_slide(
        "LineCrossDetector.update — Detect a side switch",
        ACCENT_PURPLE,
        ("Per-track state machine: stores last side seen, fires a crossing "
         "event the moment a track switches sides, and reports the "
         "direction (in vs out)."),
        ["FT-55", "FT-56", "FT-57", "FT-58"])


def ft_lcd_state():
    function_cases_slide(
        "LineCrossDetector.cleanup / get_state — Pruning & read-out",
        ACCENT_PURPLE,
        ("cleanup() drops stale per-track entries so the dict doesn't "
         "grow forever. get_state() returns the current entry/exit "
         "counts that the report consumes."),
        ["FT-59", "FT-60", "FT-61"])


# ----------------------------------------------------------------------------
# Face Recognition slides
# ----------------------------------------------------------------------------
def fr_intro():
    feature_intro_slide(
        "Face Recognition — Why It Matters",
        "25 cases  ·  tell staff and customers apart, group repeat visitors",
        ACCENT_TEAL,
        where_it_sits=(
            "Phase 3 of the pipeline. After tracking and capture, every "
            "person folder has 0..10 face crops. FR turns those into "
            "embeddings, asks FAISS 'is this any registered staff member?', "
            "and clusters the rest into unique customers."),
        failure_modes=[
            "Wrong threshold → every customer is recognised as staff (or vice "
            "versa). The Excel report still 'looks fine'.",
            "Empty staff index handled badly → run crashes on day one of "
            "deployment, before anyone is registered.",
            "Clustering merges everybody → all visitors collapse into one "
            "customer; or merges nobody → every visit is a 'new' customer.",
            "Junk-file or body-snapshot leakage → embeddings are computed on "
            "non-face data and silently corrupt the registry.",
        ],
        why_it_matters=(
            "FR is the only stage that emits the report's headline numbers — "
            "named-staff hits and customer counts. We test every entry point "
            "(extract / pick-diverse / classify / cluster) in isolation and "
            "on real face folders so a model swap or threshold tweak can't "
            "silently rewrite the metrics.")
    )


def fr_functions():
    feature_functions_slide(
        "Face Recognition — Functions Under Test",
        "Each row maps a real function to the cases that exercise it",
        ACCENT_TEAL,
        function_rows=[
            ["extract_embeddings", "5",
             "Scan a person folder, embed each face, exclude body_snapshot/junk"],
            ["pick_diverse", "6",
             "Pick mean + most-distant embeddings for the staff registry"],
            ["classify_with_faiss", "7",
             "Match against the FAISS staff index at threshold 0.80"],
            ["greedy_cluster", "7",
             "Group customers across tracks at threshold 0.60"],
        ])


def fr_extract():
    function_cases_slide(
        "extract_embeddings — Folder → unit-norm embeddings",
        ACCENT_TEAL,
        ("Walks a person folder, embeds every .jpg face crop, and returns "
         "a list of unit-length 512-d vectors. body_snapshot.jpg is "
         "excluded from the regular scan but used as a fallback when no "
         "face crop is usable."),
        ["FR-EX-01", "FR-EX-02", "FR-EX-03", "FR-EX-04", "FR-EX-05"])


def fr_pick_diverse():
    function_cases_slide(
        "pick_diverse — Mean + most-distant for the staff registry",
        ACCENT_TEAL,
        ("Given the embeddings of a registered staff member, returns the "
         "mean plus up to N most-distant samples. This is what makes the "
         "registry robust to lighting and angle changes."),
        ["FR-PD-01", "FR-PD-02", "FR-PD-03", "FR-PD-04", "FR-PD-05", "FR-PD-06"])


def fr_classify_1():
    function_cases_slide(
        "classify_with_faiss — Staff lookup",
        ACCENT_TEAL,
        ("Runs each person's embeddings through FAISS and assigns a label: "
         "staff / customer / no_face. similarity_threshold=0.80 separates "
         "staff from customer; the highest similarity wins."),
        ["FR-SM-01", "FR-SM-02", "FR-SM-03", "FR-SM-04"],
        slide_label="1 of 2")


def fr_classify_2():
    function_cases_slide(
        "classify_with_faiss — Multi-embedding & defensive paths",
        ACCENT_TEAL,
        ("These cases pin behaviour when a person has multiple embeddings, "
         "when the staff index is empty (day-1 deployment), and when the "
         "registry has fewer rows than FAISS expects."),
        ["FR-SM-05", "FR-SM-06", "FR-SM-07"],
        slide_label="2 of 2")


def fr_cluster_1():
    function_cases_slide(
        "greedy_cluster — Customer grouping",
        ACCENT_TEAL,
        ("Greedily groups non-staff persons into customer clusters using "
         "customer_cluster_threshold=0.60. The first sorted person seeds "
         "each cluster — deterministic, not globally optimal."),
        ["FR-CC-01", "FR-CC-02", "FR-CC-03", "FR-CC-04"],
        slide_label="1 of 2")


def fr_cluster_2():
    function_cases_slide(
        "greedy_cluster — Boundary & exclusion cases",
        ACCENT_TEAL,
        ("These cases cover the exact-threshold merge behaviour, the "
         "no_face exclusion, and the trivial 'all duplicates' collapse."),
        ["FR-CC-06", "FR-CC-07", "FR-CC-08"],
        slide_label="2 of 2")


# ----------------------------------------------------------------------------
# Wrap-up slides
# ----------------------------------------------------------------------------
def slide_tooling():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Tooling & Test Runner",
                     "One command runs the whole suite")
    add_rect(s, Inches(0.7), Inches(1.2), Inches(6.0), Inches(5.4),
             WHITE, line=BORDER)
    add_rect(s, Inches(0.7), Inches(1.2), Inches(6.0), Inches(0.55), NAVY)
    add_text(s, Inches(0.7), Inches(1.2), Inches(6.0), Inches(0.55),
             "  Master Runner", size=16, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.95), Inches(1.95), Inches(5.6), Inches(0.4),
             "Unit_Test_Data/run_all_and_report.py",
             size=13, bold=True, color=NAVY, font="Consolas")
    add_bullets(s, Inches(0.95), Inches(2.4), Inches(5.6), Inches(4.0),
                [
                    "Runs every test_fdNN / test_ftNN / test_fr_* script",
                    "Aggregates results into one state JSON",
                    "Renders a 4-sheet Excel report with verdicts",
                    "--fd-only / --fr-only / --ft-only filters",
                    "--from-json replays a saved run without re-execution",
                    "--skip-run rebuilds Excel from cached results.json",
                    "Per-script timeout (default 300 s, configurable)",
                ],
                size=12)
    add_rect(s, Inches(7.0), Inches(1.2), Inches(5.7), Inches(5.4),
             WHITE, line=BORDER)
    add_rect(s, Inches(7.0), Inches(1.2), Inches(5.7), Inches(0.55), ACCENT_TEAL)
    add_text(s, Inches(7.0), Inches(1.2), Inches(5.7), Inches(0.55),
             "  Output Artifacts", size=16, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    rows = [
        ["test_state_<ts>.json", "Full machine-readable run record"],
        ["FR_Thor_VideoTest_Report_<ts>.xlsx", "Summary + FD + FT + FR sheets"],
        ["Face_Tracking/data/test_ftNN/results.json", "Per-script raw outcome"],
        ["Annotation/*.mp4", "Annotated tracking videos for review"],
    ]
    add_table(s, Inches(7.2), Inches(1.95), Inches(5.3), Inches(2.6),
              ["Artifact", "Purpose"], rows,
              col_widths=[2.6, 2.7], body_size=10, header_size=11)
    add_text(s, Inches(7.2), Inches(4.7), Inches(5.3), Inches(0.4),
             "Verdicts", size=14, bold=True, color=NAVY)
    add_bullets(s, Inches(7.2), Inches(5.05), Inches(5.3), Inches(1.5),
                [
                    "PASS — assertion held",
                    "FAIL — assertion failed (with diff in notes)",
                    "SKIP — model or fixture unavailable",
                    "ERROR — script crashed or timed out",
                ],
                size=12)


def slide_environment():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Test Environment", "What the suite needs to run")
    add_text(s, Inches(0.7), Inches(1.2), Inches(12), Inches(0.4),
             "Software", size=16, bold=True, color=NAVY)
    sw_rows = [
        ["Python", "3.9+"],
        ["YOLO weights", "yolo11l.pt, yolov11s-face.pt"],
        ["InsightFace", "antelopev2 (w600k_r50.onnx)"],
        ["OSNet", "via boxmot"],
        ["FAISS", "faiss-cpu < 1.8"],
        ["NumPy", "< 2"],
    ]
    add_table(s, Inches(0.7), Inches(1.65), Inches(6.0), Inches(2.6),
              ["Component", "Version / Notes"], sw_rows,
              col_widths=[2.0, 4.0], body_size=11)
    add_text(s, Inches(7.0), Inches(1.2), Inches(6), Inches(0.4),
             "Data", size=16, bold=True, color=NAVY)
    data_rows = [
        ["data/2026-04-29/videos/", "Real CCTV clips (nested videos/videos/)"],
        ["data/2026-04-29/faces/", "Pipeline-extracted face crops"],
        ["data/2026-04-29/tracks/", "Cross-clip OSNet bank .npy files"],
        ["vector_db/", "FAISS index + staff_registry.json"],
    ]
    add_table(s, Inches(7.0), Inches(1.65), Inches(5.7), Inches(2.6),
              ["Path", "Contents"], data_rows,
              col_widths=[2.5, 3.2], body_size=11)
    add_rect(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(2.2),
             LIGHT, line=BORDER)
    add_text(s, Inches(1.0), Inches(4.65), Inches(11.5), Inches(0.4),
             "Reproducibility", size=15, bold=True, color=NAVY)
    add_text(s, Inches(1.0), Inches(5.05), Inches(11.5), Inches(1.6),
             ("• Every run records run_ts, run_tag, data_source and "
              "test_approach in the state JSON.\n"
              "• PYTHONUTF8=1 required on Windows (Unicode box characters).\n"
              "• GPU is optional — CPU works for the entire suite.\n"
              "• Saved state JSONs replay deterministically via "
              "--from-json <path>."),
             size=12, color=NAVY)


def slide_results():
    s = prs.slides.add_slide(BLANK)
    add_slide_header(s, "Results & Status",
                     f"Latest run: {STATE['run_ts']}")
    add_rect(s, Inches(0.7), Inches(1.2), Inches(12.0), Inches(1.4),
             ACCENT_GREEN)
    add_text(s, Inches(0.7), Inches(1.2), Inches(12.0), Inches(1.4),
             "ALL 130 TESTS PASS", size=40, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rows = [
        ["Face Detection",   "44", "44", "0", "0", "0"],
        ["Face Tracking",    "61", "61", "0", "0", "0"],
        ["Face Recognition", "25", "25", "0", "0", "0"],
        ["Total",           "130","130", "0", "0", "0"],
    ]
    add_table(s, Inches(0.7), Inches(2.9), Inches(12.0), Inches(2.4),
              ["Feature", "Cases", "PASS", "FAIL", "SKIP", "ERROR"], rows,
              col_widths=[3.0, 1.5, 1.5, 1.5, 1.5, 1.5],
              body_size=13, header_size=13)
    add_text(s, Inches(0.7), Inches(5.5), Inches(12), Inches(0.4),
             "What's Next", size=16, bold=True, color=NAVY)
    add_bullets(s, Inches(0.7), Inches(5.95), Inches(12), Inches(1.4),
                [
                    "Extended cases (FD-045+, FT-062+, FR-51+) are designed but not yet implemented",
                    "Wire run_all_and_report.py into CI on a small fixture clip set",
                    "Track per-test runtime to catch performance regressions early",
                    "Add a regression baseline so future model swaps surface drift",
                ],
                size=12)


def slide_closing():
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, SW, SH, NAVY)
    add_rect(s, 0, Inches(3.4), SW, Inches(0.06), ACCENT)
    add_text(s, Inches(0.7), Inches(2.4), Inches(12), Inches(1.0),
             "Thank You", size=60, bold=True, color=WHITE)
    add_text(s, Inches(0.7), Inches(3.7), Inches(12), Inches(0.6),
             "Questions & Discussion", size=24,
             color=RGBColor(0xCD, 0xDA, 0xE9))
    add_text(s, Inches(0.7), Inches(5.3), Inches(12), Inches(1.5),
             ("FR_Thor — Offline Face Recognition Pipeline\n"
              "Unit Testing Suite · Real-Footage Integration\n"
              f"130 Cases · 100% PASS · {STATE['run_date']}"),
             size=14, color=RGBColor(0xB6, 0xC4, 0xD8))


# ----------------------------------------------------------------------------
# Build deck
# ----------------------------------------------------------------------------
def main():
    # Section 1: framing
    slide_title()
    slide_agenda()
    slide_project_overview()
    slide_pipeline_arch()
    slide_why_test()
    slide_testing_approach()
    slide_how_to_read()
    slide_suite_overview()

    # Section 2: Face Detection
    fd_intro()
    fd_functions()
    fd_score_face_1()
    fd_score_face_2()
    fd_score_face_3()
    fd_check_quality()
    fd_skin()
    fd_keypoints_outer()
    fd_kps_in_bbox()
    fd_enhance_contrast()
    fd_detect_face()
    fd_strategies()
    fd_batch_capture_1()
    fd_batch_capture_2()

    # Section 3: Face Tracking
    ft_intro()
    ft_functions()
    ft_osnet_batch()
    ft_osnet_single()
    ft_active_ref()
    ft_active_update()
    ft_active_serialise()
    ft_pending()
    ft_box_distance()
    ft_match_create_1()
    ft_match_create_2()
    ft_match_create_3()
    ft_get_lost()
    ft_serialize_remove()
    ft_lcd_side()
    ft_lcd_update()
    ft_lcd_state()

    # Section 4: Face Recognition
    fr_intro()
    fr_functions()
    fr_extract()
    fr_pick_diverse()
    fr_classify_1()
    fr_classify_2()
    fr_cluster_1()
    fr_cluster_2()

    # Section 5: wrap-up
    slide_tooling()
    slide_environment()
    slide_results()
    slide_closing()

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    target = OUT_PATH
    try:
        prs.save(target)
    except PermissionError:
        # Original is open in PowerPoint — fall back to a versioned name
        i = 2
        while True:
            alt = OUT_PATH.with_name(f"{OUT_PATH.stem}_v{i}{OUT_PATH.suffix}")
            try:
                prs.save(alt)
                target = alt
                break
            except PermissionError:
                i += 1
                if i > 50:
                    raise
    print(f"Wrote {target}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
